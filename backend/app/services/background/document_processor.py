"""
Document processing service — Embeddings & Chunking.

Used by process_document_task for text extraction, chunking,
and preparation for vector embedding generation.

NOTE: Document INTAKE (upload → parse → store → index) is now
handled by the DMS ingestion pipeline (services/ingestion/).
This module handles the EMBEDDING phase that runs AFTER ingestion.
"""

import logging
from pathlib import Path
from typing import Optional, Dict, Any
import re

logger = logging.getLogger(__name__)

# Minimum text length to consider extraction successful
MIN_TEXT_LENGTH = 50


class DocumentProcessor:
    """
    Process documents to extract text for embedding generation.

    Used by process_document_task for chunking + vector embedding.
    Document intake is handled by services/ingestion/pipeline.py.
    """

    def __init__(self, enable_ocr: bool = True):
        """
        Initialize document processor.

        Args:
            enable_ocr: Enable OCR for scanned documents
        """
        self.enable_ocr = enable_ocr
        self.supported_formats = {
            "pdf": self._extract_pdf,
            "txt": self._extract_txt,
            "md": self._extract_txt,
            "docx": self._extract_docx,
            # EPUB (ebook format) — chapter-aware extraction
            "epub": self._extract_epub,
            # Image formats — visual embedding via embed_image_task is preferred;
            # these OCR fallbacks are only used when content_text is absent
            "png": self._extract_image,
            "jpg": self._extract_image,
            "jpeg": self._extract_image,
            "tiff": self._extract_image,
            "tif": self._extract_image,
            "bmp": self._extract_image,
            "gif": self._extract_image,
            "webp": self._extract_image,
        }

    # Image types that should be routed to embed_image_task, not OCR chunking
    IMAGE_TYPES = {"png", "jpg", "jpeg", "tiff", "tif", "bmp", "gif", "webp"}

    @classmethod
    def is_image_type(cls, file_type: str) -> bool:
        """True if the file type should be handled by embed_image_task (visual embeddings)."""
        return file_type.lower().lstrip(".") in cls.IMAGE_TYPES

    def process_document(self, file_path: str, file_type: str) -> Dict[str, Any]:
        """
        Process a document and extract its content.

        Args:
            file_path: Path to the document file
            file_type: File extension (pdf, txt, docx, etc.)

        Returns:
            Dictionary containing:
                - content_text: Extracted text
                - page_count: Number of pages (for PDFs)
                - word_count: Number of words
                - metadata: Additional metadata
                - ocr_performed: Whether OCR was used
        """
        # Normalize file type
        file_type = file_type.lower().lstrip(".")

        logger.info(f"Processing document: {file_path} (type: {file_type})")

        if file_type not in self.supported_formats:
            raise ValueError(f"Unsupported file type: {file_type}")

        # Extract text using appropriate method
        extractor = self.supported_formats[file_type]
        result = extractor(file_path)

        # Clean and normalize text
        content_text = self._clean_text(result["content"])

        # Calculate word count
        word_count = len(content_text.split())

        logger.info(
            f"Document processed: {word_count} words, "
            f"{result.get('page_count', 'N/A')} pages, "
            f"OCR: {result.get('ocr_performed', False)}"
        )

        return {
            "content_text": content_text,
            "page_count": result.get("page_count"),
            "word_count": word_count,
            "metadata": result.get("metadata", {}),
            "ocr_performed": result.get("ocr_performed", False),
        }

    def _extract_pdf(self, file_path: str) -> Dict[str, Any]:
        """
        Extract text from PDF file.

        Falls back to OCR if insufficient text is found.
        """
        try:
            from pypdf import PdfReader

            reader = PdfReader(file_path)
            page_count = len(reader.pages)

            # Extract text from all pages
            text_parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

            content = "\n\n".join(text_parts)

            # Extract metadata
            metadata = {}
            if reader.metadata:
                metadata = {
                    "title": reader.metadata.get("/Title"),
                    "author": reader.metadata.get("/Author"),
                    "subject": reader.metadata.get("/Subject"),
                    "creator": reader.metadata.get("/Creator"),
                }

            # Check if we have enough text, or need OCR
            if len(content) < MIN_TEXT_LENGTH and self.enable_ocr:
                logger.info(f"Insufficient text ({len(content)} chars), attempting OCR...")
                ocr_result = self._extract_with_ocr(file_path, "application/pdf")
                if ocr_result:
                    return {
                        "content": ocr_result["text"],
                        "page_count": ocr_result.get("page_count", page_count),
                        "metadata": metadata,
                        "ocr_performed": True,
                    }

            return {
                "content": content,
                "page_count": page_count,
                "metadata": metadata,
                "ocr_performed": False,
            }

        except Exception as e:
            logger.error(f"Error extracting PDF: {str(e)}")
            raise

    def _extract_image(self, file_path: str) -> Dict[str, Any]:
        """Extract text from image using OCR."""
        if not self.enable_ocr:
            raise ValueError("OCR is disabled, cannot process image files")

        # Determine MIME type from extension
        ext = Path(file_path).suffix.lower()
        mime_map = {
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".tiff": "image/tiff",
            ".tif": "image/tiff",
            ".bmp": "image/bmp",
            ".gif": "image/gif",
            ".webp": "image/webp",
        }
        mime_type = mime_map.get(ext, "image/png")

        ocr_result = self._extract_with_ocr(file_path, mime_type)

        return {
            "content": ocr_result.get("text", ""),
            "page_count": ocr_result.get("page_count", 1),
            "metadata": {},
            "ocr_performed": True,
        }

    def _extract_with_ocr(self, file_path: str, mime_type: str) -> Optional[Dict[str, Any]]:
        """
        Extract text using OCR processor.

        Args:
            file_path: Path to file
            mime_type: MIME type of file

        Returns:
            OCR result dict or None on failure
        """
        try:
            from app.services.ocr import OcrProcessor

            with OcrProcessor() as processor:
                result = processor.process_file(file_path, mime_type)
                return result

        except ImportError:
            logger.warning("OCR service not available (ocrmypdf not installed)")
            return None
        except Exception as e:
            logger.error(f"OCR failed: {e}")
            return None

    def _extract_txt(self, file_path: str) -> Dict[str, Any]:
        """Extract text from plain text file."""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()

            return {"content": content, "page_count": None, "metadata": {}}

        except UnicodeDecodeError:
            # Try with different encoding
            with open(file_path, "r", encoding="latin-1") as f:
                content = f.read()

            return {"content": content, "page_count": None, "metadata": {}}

    def _extract_docx(self, file_path: str) -> Dict[str, Any]:
        """Extract text from DOCX file."""
        try:
            from docx import Document

            doc = Document(file_path)

            # Extract text from paragraphs
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)

            content = "\n\n".join(text_parts)

            # Extract metadata
            metadata = {}
            if doc.core_properties:
                metadata = {
                    "title": doc.core_properties.title,
                    "author": doc.core_properties.author,
                    "subject": doc.core_properties.subject,
                    "keywords": doc.core_properties.keywords,
                }

            return {"content": content, "page_count": None, "metadata": metadata}

        except Exception as e:
            logger.error(f"Error extracting DOCX: {str(e)}")
            raise

    def _extract_epub(self, file_path: str) -> Dict[str, Any]:
        """
        Extract text from an EPUB file, preserving chapter structure.

        Strategy:
          1. Parse with ebooklib to get ordered document items (chapters)
          2. Skip navigation, TOC, cover, copyright items
          3. Clean each chapter's HTML with BeautifulSoup + lxml
          4. Insert double-newlines around headings to create semantic boundaries
          5. Join chapters with '\n\n' separators

        The resulting text structure helps the semantic chunker detect chapter/section
        boundaries — topics shift at headings and chapter breaks, not at character limits.

        Returns:
            content:    Full chapter-ordered plain text
            page_count: Estimated from word count (250 words/page)
            metadata:   Dublin Core title, author, subject
            ocr_performed: False (EPUBs are text-native)
        """
        try:
            import ebooklib
            from ebooklib import epub as epub_lib
            from bs4 import BeautifulSoup

            # Suppress ebooklib's noisy warnings about missing cover items
            import warnings
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                book = epub_lib.read_epub(file_path, options={"ignore_ncx": True})

            # Skip non-content documents (TOC, nav, cover, copyright)
            SKIP_HINTS = {"toc", "nav", "cover", "copyright", "title", "colophon", "halftitle"}

            chapters = []
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                name = item.get_name().lower()

                # Skip navigation/metadata documents
                if any(hint in name for hint in SKIP_HINTS):
                    continue

                html = item.get_body_content()
                if not html:
                    continue

                soup = BeautifulSoup(html, "lxml")

                # Strip noise: scripts, styles, hidden elements
                for tag in soup(["script", "style", "meta", "link"]):
                    tag.decompose()

                # Preserve heading structure as semantic boundary markers.
                # Double-newlines around headings signal topic shifts to the
                # semantic chunker, keeping chapters from bleeding into each other.
                for heading in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
                    heading.string = f"\n\n{heading.get_text(strip=True)}\n\n"

                text = soup.get_text(separator=" ")
                # Collapse whitespace runs but preserve paragraph breaks
                text = re.sub(r"[ \t]+", " ", text)
                text = re.sub(r"\n{3,}", "\n\n", text).strip()

                if len(text) >= MIN_TEXT_LENGTH:
                    chapters.append(text)

            if not chapters:
                raise ValueError(f"No readable content extracted from EPUB: {file_path}")

            # Join with section separators
            full_text = "\n\n".join(chapters)

            # Estimate page count (250 words per page is standard)
            word_count = len(full_text.split())
            estimated_pages = max(1, word_count // 250)

            # Extract Dublin Core metadata
            metadata: Dict[str, Any] = {}
            dc_title = book.get_metadata("DC", "title")
            if dc_title:
                metadata["title"] = dc_title[0][0]
            dc_creator = book.get_metadata("DC", "creator")
            if dc_creator:
                metadata["author"] = dc_creator[0][0]
            dc_subject = book.get_metadata("DC", "subject")
            if dc_subject:
                metadata["subject"] = dc_subject[0][0]

            logger.info(
                f"EPUB extracted: {len(chapters)} chapters, "
                f"{word_count:,} words, ~{estimated_pages} pages"
            )

            return {
                "content": full_text,
                "page_count": estimated_pages,
                "metadata": metadata,
                "ocr_performed": False,
            }

        except Exception as e:
            logger.error(f"Error extracting EPUB {file_path}: {e}")
            raise

    def _clean_text(self, text: str) -> str:
        """
        Clean and normalize extracted text.

        IMPORTANT: Only collapses horizontal whitespace (spaces and tabs).
        Newlines are deliberately preserved — they carry paragraph and chapter
        structure that the semantic chunker uses to find topic boundaries.
        Collapsing \\n to spaces (the old behavior) destroyed all structural
        markers before the chunker ever saw the text, causing 98% safeguard rate.

        Args:
            text: Raw extracted text

        Returns:
            Cleaned text with paragraph structure intact
        """
        # Remove NUL bytes — PostgreSQL TEXT columns reject 0x00 characters
        text = text.replace("\x00", "")

        # Collapse horizontal whitespace only (spaces and tabs).
        # DO NOT use r"\s+" here — \s matches \n which destroys paragraph breaks.
        text = re.sub(r"[ \t]+", " ", text)

        # Normalize excessive newlines (3+ → 2) while preserving paragraph breaks
        text = re.sub(r"\n{3,}", "\n\n", text)

        # Strip leading/trailing whitespace
        text = text.strip()

        return text

    def iter_chunks(
        self,
        text: str,
        document_id: str,
        chunk_size: int = 1000,
        overlap: int = 200,
    ):
        """
        Yield chunks one at a time for memory-safe streaming ingestion.

        STREAMING INVARIANT: At no point should the full chunk list exist in memory.
        This generator yields chunks lazily, allowing batch processing with gc.collect().

        Args:
            text: Text to chunk
            document_id: Document ID for deterministic chunk hashing
            chunk_size: Target chunk size in characters
            overlap: Overlap between chunks in characters

        Yields:
            Dict with content, chunk_index, start_char, end_char, chunk_id
        """
        import uuid

        # Safety limits
        MAX_CHUNKS = 50000  # Absolute maximum chunks per document
        MIN_ADVANCE = max(1, chunk_size - overlap)  # Minimum progress per iteration

        start = 0
        chunk_index = 0
        last_start = -1  # For infinite loop detection

        while start < len(text) and chunk_index < MAX_CHUNKS:
            # Infinite loop detection
            if start == last_start:
                logger.error(f"Chunker stuck at position {start}, forcing advance")
                start += MIN_ADVANCE
                continue
            last_start = start

            end = start + chunk_size

            # Find the nearest sentence boundary
            if end < len(text):
                best_boundary = None
                for punct in [". ", "! ", "? ", "\n\n"]:
                    punct_pos = text.rfind(punct, start, end)
                    if punct_pos != -1:
                        # Take the furthest boundary to maximize chunk size
                        if best_boundary is None or punct_pos > best_boundary:
                            best_boundary = punct_pos + len(punct)

                if best_boundary is not None:
                    end = best_boundary

            # Clamp end to text length
            end = min(end, len(text))

            chunk_text = text[start:end].strip()

            if chunk_text:
                # Deterministic chunk ID for idempotent re-ingestion
                chunk_id = str(
                    uuid.uuid5(uuid.NAMESPACE_OID, f"{document_id}:{start}:{chunk_size}:{overlap}")
                )

                yield {
                    "content": chunk_text,
                    "chunk_index": chunk_index,
                    "start_char": start,
                    "end_char": end,
                    "chunk_id": chunk_id,
                }
                chunk_index += 1

            # CRITICAL: Guarantee forward progress
            # Move to next position with overlap, but ensure minimum advance
            next_start = end - overlap
            if next_start <= start:
                # Overlap would cause no progress - force advance
                next_start = start + MIN_ADVANCE
            start = next_start

    def chunk_text(
        self, text: str, chunk_size: int = 1000, overlap: int = 200
    ) -> list[Dict[str, Any]]:
        """
        Split text into overlapping chunks for embedding.

        COMPATIBILITY WRAPPER: For streaming ingestion, use iter_chunks() instead.

        Args:
            text: Text to chunk
            chunk_size: Target chunk size in characters
            overlap: Overlap between chunks in characters

        Returns:
            List of chunk dictionaries with content and metadata
        """
        # Use generator but collect for backwards compatibility
        chunks = list(
            self.iter_chunks(
                text=text,
                document_id="legacy",  # No document ID in legacy calls
                chunk_size=chunk_size,
                overlap=overlap,
            )
        )

        logger.info(f"Created {len(chunks)} chunks from text")
        return chunks
